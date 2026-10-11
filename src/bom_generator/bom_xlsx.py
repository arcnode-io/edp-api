"""Bom → xlsx bytes, one row per line item."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from src.bom_generator.bom_models import Bom


# Column order matches BomLineItem field order; metadata cols last so a
# consumer eyeballing the sheet sees procurement essentials first.
_XLSX_COLUMNS: tuple[tuple[str, str], ...] = (
    ("part_number", "Part Number"),
    ("vendor", "Vendor"),
    ("description", "Description"),
    ("qty", "Qty"),
    ("procurement_path", "Procurement Path"),
    ("unit_cost_usd", "Unit Cost (USD)"),
    ("lead_time_weeks", "Lead Time (weeks)"),
    ("datasheet_url", "Datasheet"),
    ("install_video_url", "Install Video"),
    ("ndaa_compliant", "NDAA"),
    ("taa_compliant", "TAA"),
    # Track-B derived columns from `offers`: cheapest live price + source.
    # Empty when enrichment didn't run for this row (no distributor returned
    # a non-error offer). Full per-distributor breakdown lives in the json.
    ("__live_cheapest", "Live Cheapest (USD)"),
    ("__live_source", "Live Source"),
    ("price_change_pct_7d", "Δ vs 7d ago (%)"),
    ("material", "Material"),
    ("finish", "Finish"),
    ("drawing_ref", "Drawing Ref"),
    ("drawing_url", "Drawing URL"),
)


def _cheapest_offer(offers: list) -> tuple[float | None, str | None]:  # type: ignore[type-arg]
    """Lowest non-error offer's (unit_cost_usd, distributor). None when no priced offers."""
    priced = [o for o in offers if o.error is None and o.unit_cost_usd is not None]
    if not priced:
        return None, None
    best = min(priced, key=lambda o: o.unit_cost_usd)
    return best.unit_cost_usd, best.distributor


def serialize_bom_xlsx(bom: Bom) -> bytes:
    """Serialize a Bom to xlsx bytes for S3 upload.

    One header row + one data row per BomLineItem. Header row is bold.
    Empty cells for fields that don't apply to a given procurement_path.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"

    bold = Font(bold=True)
    for col_idx, (_field, label) in enumerate(_XLSX_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = bold

    for row_idx, item in enumerate(bom.line_items, start=2):
        cheapest_price, cheapest_source = _cheapest_offer(item.offers)
        for col_idx, (field, _label) in enumerate(_XLSX_COLUMNS, start=1):
            if field == "__live_cheapest":
                value = cheapest_price
            elif field == "__live_source":
                value = cheapest_source
            else:
                value = getattr(item, field)
            # Reason: openpyxl writes StrEnum as the enum object, not its value.
            if hasattr(value, "value"):
                value = value.value
            ws.cell(row=row_idx, column=col_idx, value=value)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
