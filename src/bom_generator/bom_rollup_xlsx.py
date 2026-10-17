"""Purchase roll-up sheet: one row per part the customer buys, summed."""

from typing import Final

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.bom_generator.bom_models import Bom, BomLineItem
from src.bom_generator.bom_pricing import PURCHASED

_BOLD: Final[Font] = Font(bold=True)
_HEAD_FILL: Final[PatternFill] = PatternFill("solid", fgColor="D9D9D9")


def rollup_sheet(ws: Worksheet, bom: Bom) -> None:
    """One row per part the customer buys, quantities summed across sections."""
    headers = [
        "Manufacturer",
        "Part number",
        "Qty",
        "UoM",
        "Description",
        "Sections",
        "Unit cost (USD)",
        "Ext. cost (USD)",
        "Lead time (wk)",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font, cell.fill = _BOLD, _HEAD_FILL
    # Reason: key on equipment_id too, so two engineered-to-order parts that
    # share a vendor placeholder stay separate rows.
    parts: dict[tuple[str, str, str | None], list[BomLineItem]] = {}
    for li in bom.line_items:
        if li.procurement_path in PURCHASED:
            parts.setdefault((li.vendor, li.part_number, li.equipment_id), []).append(
                li
            )
    for (vendor, part, _), lines in parts.items():
        qty = sum(li.qty for li in lines)
        unit = lines[0].unit_cost_usd
        lead = lines[0].lead_time_weeks
        ws.append(
            [
                vendor,
                part,
                qty,
                lines[0].uom,
                lines[0].description,
                ", ".join(sorted({li.section.value for li in lines})),
                "Unpriced" if unit is None else unit,
                "Unpriced" if unit is None else unit * qty,
                "RFQ" if lead is None else lead,
            ]
        )
    for i, width in enumerate((24, 30, 6, 6, 60, 30, 14, 14, 10), start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
