"""BOM workbook: order header, sectioned lines with no empty cell, purchase roll-up."""

from datetime import UTC, date, datetime
from io import BytesIO
from uuid import UUID

from openpyxl import load_workbook

from src.bom_generator.bom_models import (
    Bom,
    BomLineItem,
    BomOrder,
    BomSection,
    ProcurementPath,
)
from src.bom_generator.bom_site_lines import site_lines
from src.bom_generator.bom_xlsx import serialize_bom_xlsx
from src.bom_generator.pricing_models import PriceSource

_ORDER = BomOrder(
    customer="acme",
    site="brookside dc-1",
    deployment_context="commercial",
    gpu_variant="h100_sxm",
    gpus_requested=8,
    gpus_delivered=56,
    bess_capacity_mwh=1.0,
    bess_coupling="ac_coupled",
)


def _line(
    section: BomSection, part: str, qty: int, path: ProcurementPath
) -> BomLineItem:
    return BomLineItem(
        section=section,
        ref_des="C1",
        part_number=part,
        vendor="Acme",
        description=part,
        qty=qty,
        procurement_path=path,
    )


def _bom() -> Bom:
    node = _line(BomSection.COMPUTE, "SYS-421GE-TNHR2-LCC", 7, ProcurementPath.CATALOG)
    lines = [
        node.model_copy(
            update={
                "equipment_id": "CMP-NODE-002",
                "fab_tier": "dod_eligible",
                "ndaa_compliant": True,
                "taa_compliant": True,
                "datasheet_url": "https://example.com/ds",
                "unit_cost_usd": 300_000.0,
                "price_source": PriceSource.ESTIMATE,
                "price_basis": "reseller list",
                "price_as_of": date(2026, 10, 17),
            }
        ),
        _line(
            BomSection.COMPUTE, "ARC-PLT-CG-001", 1, ProcurementPath.CUSTOM_FABRICATION
        ),
        _line(BomSection.GRID, "ARC-PLT-CG-001", 1, ProcurementPath.CUSTOM_FABRICATION),
        _line(
            BomSection.GRID,
            "To spec GRD-LVS-001",
            1,
            ProcurementPath.ENGINEERED_TO_ORDER,
        ),
        *site_lines(_ORDER, compute_qty=1),
        _line(BomSection.FIELD, "Cat6 STP", 17, ProcurementPath.FIELD_SUPPLIED),
    ]
    return Bom(
        deployment_id=UUID("00000000-0000-0000-0000-000000000045"),
        profile="commercial_ac",
        manifest_version="0.0.0-test",
        generated_at=datetime.now(UTC),
        compute_container_qty=1,
        grid_container_qty=1,
        order=_ORDER,
        line_items=[li.model_copy(update={"item": i}) for i, li in enumerate(lines, 1)],
    )


def _sheets() -> tuple[list[tuple], list[tuple]]:
    wb = load_workbook(BytesIO(serialize_bom_xlsx(_bom())))
    bom = list(wb["BOM"].iter_rows(values_only=True))
    rollup = list(wb["Purchase roll-up"].iter_rows(values_only=True))
    return bom, rollup


def test_header_states_the_order_as_placed_and_as_delivered() -> None:
    # Act
    bom, _ = _sheets()

    # Assert
    header = " | ".join(str(c) for r in bom[:12] for c in r if c is not None)
    assert "acme" in header
    assert "brookside dc-1" in header
    assert "8 x H100 SXM requested → 1 compute container, 56 GPUs" in header
    assert "1 MWh" in header


def test_every_cell_of_every_line_is_filled() -> None:
    # Act
    bom, _ = _sheets()

    # Assert — line rows start with their item number
    lines = [r for r in bom if isinstance(r[0], int)]
    empty = [(r[0], i) for r in lines for i, c in enumerate(r) if c in (None, "")]
    assert (len(lines), empty) == (len(_bom().line_items), [])


def test_rollup_totals_each_part_the_customer_buys_across_sections() -> None:
    # Act
    _, rollup = _sheets()

    # Assert — CG plate on compute and grid sums to 2; customer gear isn't bought
    qty = {r[1]: r[2] for r in rollup if isinstance(r[2], int)}
    assert qty["ARC-PLT-CG-001"] == 2
    assert "Customer-selected" not in qty


def test_lines_show_their_price_basis_and_the_order_a_budgetary_total() -> None:
    # Act
    bom, _ = _sheets()

    # Assert — 7 nodes at a ±40% estimate; everything else is unpriced
    cells = [c for r in bom for c in r if isinstance(c, str)]
    total = next(r for r in bom if r[1] == "BUDGET TOTAL (budgetary)")
    assert "Estimate 2026-10-17 ±40%: reseller list" in cells
    assert (total[11], total[12]) == (2_100_000, "Range $1,260,000 to $2,940,000")
