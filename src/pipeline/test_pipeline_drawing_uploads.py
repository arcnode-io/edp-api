"""The pipeline uploads real engineering drawings and schedules."""

import json

from src.bom_generator.manifest_service import ManifestService
from src.module_resolver.module_resolver_service import ModuleResolverService
from src.pipeline.artifact_urls import build_artifact_urls_from_resolved
from src.shared.enums import (
    GpuVariant,
)
from src.shared.schemas.artifact import ArtifactKind
from src.pipeline.pipeline_test_support import (
    DEPLOYMENT_ID,
    _RecordingClient,
    _commercial_ac_manifest,
    _payload,
    _build_pipeline,
)


def test_sld_engineering_uploads_real_dxf_and_pdf() -> None:
    """SLD dxf URL gets parseable DXF bytes + sld pdf URL gets %PDF- magic bytes."""
    # Arrange
    import io as _io

    import ezdxf as _ezdxf

    client = _RecordingClient(_commercial_ac_manifest())
    pipeline = _build_pipeline(client)
    payload = _payload()
    resolution = ModuleResolverService().resolve(payload)
    urls = build_artifact_urls_from_resolved(
        DEPLOYMENT_ID,
        ManifestService(manifest=_commercial_ac_manifest()).resolve(
            "commercial_ac", gpu_variant=GpuVariant.B200
        ),
    )

    # Act
    pipeline.run(
        payload=payload,
        resolution=resolution,
        urls=urls,
        manifest=_commercial_ac_manifest(),
    )

    # Assert — both URLs got real bytes, not stubs
    dxf_url = next(
        u.url for u in urls if u.kind == ArtifactKind.SLD and u.format == "dxf"
    )
    pdf_url = next(
        u.url for u in urls if u.kind == ArtifactKind.SLD and u.format == "pdf"
    )
    # DXF round-trips through ezdxf — proves real DXF.
    _ezdxf.read(_io.StringIO(client.uploads[dxf_url].decode("utf-8")))
    # PDF starts with the %PDF- magic.
    assert client.uploads[pdf_url].startswith(b"%PDF-")


def test_cable_hose_schedule_uploads_real_json_and_xlsx() -> None:
    """Cable+hose json URL gets a CableHoseSchedule shape + xlsx round-trips."""
    # Arrange
    import io as _io

    from openpyxl import load_workbook as _load_wb

    client = _RecordingClient(_commercial_ac_manifest())
    pipeline = _build_pipeline(client)
    payload = _payload()
    resolution = ModuleResolverService().resolve(payload)
    urls = build_artifact_urls_from_resolved(
        DEPLOYMENT_ID,
        ManifestService(manifest=_commercial_ac_manifest()).resolve(
            "commercial_ac", gpu_variant=GpuVariant.B200
        ),
    )

    # Act
    pipeline.run(
        payload=payload,
        resolution=resolution,
        urls=urls,
        manifest=_commercial_ac_manifest(),
    )

    # Assert — json deserializes with `cables` + `hoses` keys.
    json_url = next(
        u.url
        for u in urls
        if u.kind == ArtifactKind.CABLE_HOSE_SCHEDULE and u.format == "json"
    )
    body = json.loads(client.uploads[json_url])
    assert "cables" in body
    assert "hoses" in body

    # Assert — xlsx round-trips with Cables + Hoses sheets.
    xlsx_url = next(
        u.url
        for u in urls
        if u.kind == ArtifactKind.CABLE_HOSE_SCHEDULE and u.format == "xlsx"
    )
    wb = _load_wb(_io.BytesIO(client.uploads[xlsx_url]))
    assert "Cables" in wb.sheetnames
    assert "Hoses" in wb.sheetnames


def test_comms_diagram_uploads_real_dxf_and_pdf() -> None:
    """Comms diagram dxf URL gets parseable DXF + pdf URL gets %PDF- magic bytes."""
    # Arrange
    import io as _io

    import ezdxf as _ezdxf

    client = _RecordingClient(_commercial_ac_manifest())
    pipeline = _build_pipeline(client)
    payload = _payload()
    resolution = ModuleResolverService().resolve(payload)
    urls = build_artifact_urls_from_resolved(
        DEPLOYMENT_ID,
        ManifestService(manifest=_commercial_ac_manifest()).resolve(
            "commercial_ac", gpu_variant=GpuVariant.B200
        ),
    )

    # Act
    pipeline.run(
        payload=payload,
        resolution=resolution,
        urls=urls,
        manifest=_commercial_ac_manifest(),
    )

    # Assert
    dxf_url = next(
        u.url
        for u in urls
        if u.kind == ArtifactKind.COMMS_DIAGRAM and u.format == "dxf"
    )
    pdf_url = next(
        u.url
        for u in urls
        if u.kind == ArtifactKind.COMMS_DIAGRAM and u.format == "pdf"
    )
    _ezdxf.read(_io.StringIO(client.uploads[dxf_url].decode("utf-8")))
    assert client.uploads[pdf_url].startswith(b"%PDF-")


def test_pid_cooling_uploads_real_dxf_and_pdf() -> None:
    """P&ID dxf URL gets parseable DXF + pdf URL gets %PDF- magic bytes."""
    # Arrange
    import io as _io

    import ezdxf as _ezdxf

    client = _RecordingClient(_commercial_ac_manifest())
    pipeline = _build_pipeline(client)
    payload = _payload()
    resolution = ModuleResolverService().resolve(payload)
    urls = build_artifact_urls_from_resolved(
        DEPLOYMENT_ID,
        ManifestService(manifest=_commercial_ac_manifest()).resolve(
            "commercial_ac", gpu_variant=GpuVariant.B200
        ),
    )

    # Act
    pipeline.run(
        payload=payload,
        resolution=resolution,
        urls=urls,
        manifest=_commercial_ac_manifest(),
    )

    # Assert
    dxf_url = next(
        u.url for u in urls if u.kind == ArtifactKind.PID_COOLING and u.format == "dxf"
    )
    pdf_url = next(
        u.url for u in urls if u.kind == ArtifactKind.PID_COOLING and u.format == "pdf"
    )
    _ezdxf.read(_io.StringIO(client.uploads[dxf_url].decode("utf-8")))
    assert client.uploads[pdf_url].startswith(b"%PDF-")
